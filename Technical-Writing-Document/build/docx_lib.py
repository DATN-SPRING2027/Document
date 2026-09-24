"""Helpers that build a capstone report on top of the sample .docx (styles, cover logo, numbering are reused)."""
import copy
import os
import re

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLE = os.path.join(os.path.dirname(HERE), "Sample-Report", "1614_DN_SE_01_Report1_Project_Introduction.docx")

HDR_FILL = "FFE8E1"      # table header fill used by the sample
HDR_COLOR = "6E2500"     # table header text colour used by the sample
GROUP_FILL = "F2F2F2"
PAGE_TEXT_WIDTH_CM = 15.9

_TOKEN = re.compile(r"(\*\*.+?\*\*|\[\[.+?\]\]|`.+?`)")


class Report:
    def __init__(self, cover_title, cover_date, project_line):
        self.doc = docx.Document(SAMPLE)
        self.fig_no = 0
        body = self.doc.element.body
        children = list(body)
        sect = children[-1]
        assert sect.tag == qn("w:sectPr")
        # keep the cover (body items 0..20), drop the sample's TOC and content
        for el in children[21:-1]:
            body.remove(el)
        paras = self.doc.paragraphs
        self._set_text(paras[10], cover_title)
        self._set_text(paras[20], cover_date)
        if "[" in cover_date:
            paras[20].runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
        self._fill_cover_project(paras[12], project_line)
        self._bullet_pPr = self._find_bullet_template()

    # ---------- low level ----------
    @staticmethod
    def _set_text(p, text):
        runs = p.runs
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)

    @staticmethod
    def _fill_cover_project(p, text):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.runs[0] if p.runs else p.add_run()
        r.text = text
        r.bold = True
        r.font.size = Pt(18)
        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    def _find_bullet_template(self):
        src = docx.Document(SAMPLE)
        for p in src.paragraphs:
            if p.text.startswith("Project name"):
                return copy.deepcopy(p._p.pPr)
        raise RuntimeError("bullet template not found in sample")

    def _runs(self, p, text, size=None, bold=None, italic=None, color=None):
        for part in _TOKEN.split(text):
            if not part:
                continue
            hl = strong = code = False
            if part.startswith("**") and part.endswith("**"):
                part, strong = part[2:-2], True
            elif part.startswith("[[") and part.endswith("]]"):
                part, hl = "[" + part[2:-2] + "]", True
            elif part.startswith("`") and part.endswith("`"):
                part, code = part[1:-1], True
            r = p.add_run(part)
            if size:
                r.font.size = Pt(size)
            if bold or strong:
                r.bold = True
            if italic:
                r.italic = True
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
            if hl:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if code:
                r.font.name = "Consolas"
                r._r.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
                if size is None:
                    r.font.size = Pt(10)
        return p

    # ---------- blocks ----------
    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h1(self, text, page_break=False):
        p = self.doc.add_paragraph(style="Heading 1")
        p.paragraph_format.page_break_before = page_break
        p.paragraph_format.space_after = Pt(6)
        self._runs(p, text)
        return p

    def h2(self, text):
        p = self.doc.add_paragraph(style="Heading 2")
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        self._runs(p, text)
        return p

    def h3(self, text):
        p = self.doc.add_paragraph(style="Heading 3")
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        self._runs(p, text)
        return p

    def para(self, text, align=None, size=None, italic=None, after=6, before=0, indent_cm=None, keep_next=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.line_spacing = 1.1
        if indent_cm:
            p.paragraph_format.left_indent = Cm(indent_cm)
        if align == "justify":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = keep_next
        self._runs(p, text, size=size, italic=italic)
        return p

    def label(self, text):
        """Bold lead-in line such as 'Key user groups:'."""
        return self.para(f"**{text}**", after=3, before=4, keep_next=True)

    def bullets(self, items, level=0, after=2):
        for it in items:
            p = self.doc.add_paragraph()
            old = p._p.pPr
            if old is not None:
                p._p.remove(old)
            p._p.insert(0, copy.deepcopy(self._bullet_pPr))
            pf = p.paragraph_format
            pf.space_after = Pt(after)
            pf.space_before = Pt(0)
            pf.line_spacing = 1.1
            if level:
                pf.left_indent = Cm(1.6 + 0.6 * (level - 1))
            self._runs(p, it)

    def figure(self, path, width_cm, caption):
        self.fig_no += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(path, width=Cm(width_cm))
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(10)
        self._runs(c, f"Figure {self.fig_no}. {caption}", size=10, italic=True)

    def note(self, text):
        self.para(text, size=9.5, italic=True, after=8)

    # ---------- table ----------
    def table(self, headers, rows, widths_cm, size=10, group_rows=(), center_cols=(), total_row=False,
              header=True):
        ncol = len(widths_cm)
        t = self.doc.add_table(rows=0, cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        tblPr = t._tbl.tblPr
        for el in tblPr.findall(qn("w:tblStyle")):
            tblPr.remove(el)
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:type"), "dxa")
        tblW.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), "000000")
            borders.append(b)
        tblPr.append(borders)
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        mar = OxmlElement("w:tblCellMar")
        for side, w in (("top", 30), ("left", 90), ("bottom", 30), ("right", 90)):
            m = OxmlElement(f"w:{side}")
            m.set(qn("w:w"), str(w))
            m.set(qn("w:type"), "dxa")
            mar.append(m)
        tblPr.append(mar)

        def add_row(values, kind):
            row = t.add_row()
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trPr.append(cs)
            if kind == "head":
                th = OxmlElement("w:tblHeader")
                trPr.append(th)
            for i, val in enumerate(values):
                cell = row.cells[i]
                cell.width = Cm(widths_cm[i])
                tcPr = cell._tc.get_or_add_tcPr()
                fill = HDR_FILL if kind == "head" else (GROUP_FILL if kind in ("group", "total") else None)
                if fill:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), fill)
                    tcPr.append(shd)
                lines = str(val).split("\n") if val != "" else [""]
                for j, line in enumerate(lines):
                    p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                    pf = p.paragraph_format
                    pf.space_before = Pt(2)
                    pf.space_after = Pt(2)
                    pf.line_spacing = 1.05
                    if kind == "head" or i in center_cols:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    self._runs(p, line, size=size, bold=(kind in ("head", "group", "total")) or None,
                               italic=(kind == "group") or None,
                               color=HDR_COLOR if kind == "head" else None)
            return row

        if header and headers:
            add_row(headers, "head")
        for ri, r in enumerate(rows):
            kind = "group" if ri in group_rows else ("total" if (total_row and ri == len(rows) - 1) else "body")
            add_row(r, kind)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        return t

    # ---------- front matter ----------
    def toc(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.keep_with_next = True
        r = p.add_run("Table of Contents")
        r.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(0x2E, 0x75, 0xB5)
        fp = self.doc.add_paragraph()
        run = fp.add_run()
        for kind, text in (("begin", None), (None, ' TOC \\h \\u \\z \\t "Heading 1,1,Heading 2,2,Heading 3,3," '),
                           ("separate", None), (None, "Right-click and choose Update Field to refresh the table of contents."),
                           ("end", None)):
            if kind:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                run._r.append(fc)
            elif text.startswith(" TOC"):
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = text
                run._r.append(it)
            else:
                tt = OxmlElement("w:t")
                tt.text = text
                run._r.append(tt)

    def footer(self, text):
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        fp = sec.footer.paragraphs[0]
        for r in list(fp.runs):
            r._r.getparent().remove(r._r)
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        tabs = fp.paragraph_format.tab_stops
        from docx.enum.text import WD_TAB_ALIGNMENT
        tabs.add_tab_stop(Cm(PAGE_TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
        r = fp.add_run(text + "\tPage ")
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        r2 = fp.add_run()
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        for kind, text in (("begin", None), ("instr", " PAGE "), ("separate", None), ("t", "1"), ("end", None)):
            if kind == "instr":
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = text
                r2._r.append(it)
            elif kind == "t":
                tt = OxmlElement("w:t")
                tt.text = text
                r2._r.append(tt)
            else:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                r2._r.append(fc)

    def record_of_changes(self, rows, blank_rows=6):
        self.h1("I. Record of Changes")
        data = [list(r) for r in rows] + [["", "", "", ""] for _ in range(blank_rows)]
        self.table(["Date", "A*, M, D", "In charge", "Change Description"], data, [2.4, 2.0, 3.1, 8.4],
                   center_cols=(0, 1))
        self.para("*A - Added  M - Modified  D - Deleted", size=10, after=6)

    def save(self, path, title):
        cp = self.doc.core_properties
        cp.title = title
        cp.author = "Continuum AI Team"
        cp.last_modified_by = "Continuum AI Team"
        self.doc.save(path)


def refresh_with_word(path, pdf_path=None):
    """Open the file in Microsoft Word, refresh TOC/fields (real page numbers), save, optionally export PDF."""
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(os.path.abspath(path))
        d.Repaginate()
        for i in range(d.TablesOfContents.Count):
            d.TablesOfContents(i + 1).Update()
        d.Fields.Update()
        d.Repaginate()
        for i in range(d.TablesOfContents.Count):
            d.TablesOfContents(i + 1).Update()
        pages = d.ComputeStatistics(2)
        d.Save()
        if pdf_path:
            d.ExportAsFixedFormat(os.path.abspath(pdf_path), 17)
        d.Close(False)
        return pages
    finally:
        word.Quit()
